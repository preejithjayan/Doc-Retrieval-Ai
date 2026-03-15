from dataclasses import dataclass
from pathlib import Path

from services.chunking import normalize_text
from services.ocr_service import OCRService


@dataclass
class ExtractionResult:
    raw_text: str
    normalized_text: str
    used_ocr: bool = False
    ocr_confidence: float | None = None
    ocr_payload: dict | None = None
    file_type: str | None = None


class DocumentParserService:
    IMAGE_EXTENSIONS = {'.png', '.jpg', '.jpeg', '.bmp', '.tiff'}

    def __init__(self):
        self.ocr_service = OCRService()

    def detect_file_type(self, file_path):
        suffix = Path(file_path).suffix.lower()
        if suffix in self.IMAGE_EXTENSIONS:
            return 'image'
        if suffix == '.pdf':
            return 'pdf'
        if suffix == '.docx':
            return 'docx'
        if suffix == '.txt':
            return 'txt'
        return 'unknown'

    def extract(self, document):
        file_path = Path(document.file_path.path)
        file_type = self.detect_file_type(file_path)

        if file_type == 'pdf':
            return self._extract_pdf(file_path)
        if file_type == 'docx':
            return self._extract_docx(file_path)
        if file_type == 'txt':
            return self._extract_text(file_path)
        if file_type == 'image':
            return self._extract_image(file_path)
        raise ValueError(f'Unsupported document type: {file_path.suffix}')

    def _extract_text(self, file_path):
        raw_text = file_path.read_text(encoding='utf-8', errors='ignore')
        return ExtractionResult(raw_text=raw_text, normalized_text=normalize_text(raw_text), file_type='txt')

    def _extract_docx(self, file_path):
        from docx import Document as WordDocument

        document = WordDocument(str(file_path))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(' | '.join(cells))
        raw_text = "\n".join(blocks)
        return ExtractionResult(raw_text=raw_text, normalized_text=normalize_text(raw_text), file_type='docx')

    def _extract_pdf(self, file_path):
        from pypdf import PdfReader

        reader = PdfReader(str(file_path))
        pages = [page.extract_text() or '' for page in reader.pages]
        raw_text = "\n".join(pages)
        normalized = normalize_text(raw_text)
        if normalized:
            return ExtractionResult(raw_text=raw_text, normalized_text=normalized, file_type='pdf')
        return self._extract_scanned_pdf(file_path)

    def _extract_scanned_pdf(self, file_path):
        from pdf2image import convert_from_path

        pages = convert_from_path(str(file_path))
        page_text = []
        confidence_values = []
        payload = {'pages': []}
        for index, page in enumerate(pages, start=1):
            temp_image = file_path.with_name(f'{file_path.stem}_page_{index}.png')
            page.save(temp_image, 'PNG')
            extraction = self.ocr_service.extract_from_image(temp_image)
            temp_image.unlink(missing_ok=True)
            page_text.append(extraction.text)
            confidence_values.append(extraction.confidence)
            payload['pages'].append({'page': index, **extraction.structured_output})
        joined = "\n".join(page_text)
        average = round(sum(confidence_values) / len(confidence_values), 4) if confidence_values else 0.0
        return ExtractionResult(
            raw_text=joined,
            normalized_text=normalize_text(joined),
            used_ocr=True,
            ocr_confidence=average,
            ocr_payload=payload,
            file_type='pdf',
        )

    def _extract_image(self, file_path):
        extraction = self.ocr_service.extract_from_image(file_path)
        return ExtractionResult(
            raw_text=extraction.text,
            normalized_text=normalize_text(extraction.text),
            used_ocr=True,
            ocr_confidence=extraction.confidence,
            ocr_payload=extraction.structured_output,
            file_type='image',
        )
